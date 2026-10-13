from datetime import date
from openpyxl import load_workbook
import pytest

from services.api.collaboratemd import create_collaboratemd_workbook


def payload(**overrides):
    return {"encounter_id": "SYNTH-XLSX", "patient_reference": "=HYPERLINK(\"https://example.invalid\")",
            "snapshot_id": "snapshot-1", "cms1500_lines": [{
                "dos": date(2026, 9, 22), "pos": "11", "cpt": "00100",
                "justification": "=1+1", "mod1": "-NU",
            }], "diagnoses": [], **overrides}


def test_literal_text_identifiers_and_no_invented_defaults():
    stream = create_collaboratemd_workbook(payload())
    assert stream.tell() == 0
    wb = load_workbook(stream)
    assert wb.sheetnames == ["CMS-1500 (Professional)"]
    ws = wb.active
    assert ws['F7'].value == '00100'
    assert ws['F7'].number_format == '@'
    assert ws['L7'].value == '=1+1'
    assert ws['L7'].data_type == 's'
    assert ws['G7'].value == '-NU'
    assert ws['J7'].value is None  # Units have not been reviewed.
    assert ws['I7'].value is None  # No invented diagnosis pointer.
    assert ws['D7'].value is None  # No invented primary diagnosis.
    assert ws['B7'].value.date() == date(2026, 9, 22)
    assert ws.freeze_panes == 'D7'
    assert ws.auto_filter.ref == 'A6:L7'
    assert all(cell.data_type != 'f' for row in ws for cell in row)


def test_generation_does_not_create_temporary_worksheet_files(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError('Worksheet data must not be written to a temporary file.')
    monkeypatch.setattr('openpyxl.worksheet._writer.NamedTemporaryFile', forbidden)
    wb = load_workbook(create_collaboratemd_workbook(payload(is_facility=True)))
    assert len(wb.worksheets) == 2


@pytest.mark.parametrize('text', ['bad\x00text', 'x' * 32768, 'x' * 2000],
                         ids=['control-character', 'cell-length', 'row-height'])
def test_unrepresentable_evidence_fails_instead_of_truncating(text):
    data = payload()
    data['cms1500_lines'][0]['justification'] = text
    with pytest.raises(ValueError):
        create_collaboratemd_workbook(data)
