import hashlib
import hmac
import json
from io import BytesIO
from openpyxl import load_workbook
from pathlib import Path
import pytest
from sqlalchemy import select
from services.api.config import settings
from services.api.db.models import CaseSnapshot, Evidence, AuditEvent


async def prepared_case(client, seed_data):
    headers = {'Authorization': 'Bearer ' + seed_data['tokens']['reviewer_org1']}
    created = await client.post('/api/cases', headers=headers, json={
        'patient_reference': 'SYNTH-INTEGRITY', 'service_date': '2026-09-22',
        'coding_domain': 'emergency_department', 'place_of_service': '23',
    })
    assert created.status_code == 201
    case_id = created.json()['id']
    source = Path('tests/fixtures/synthetic-er-record.txt').read_bytes()
    uploaded = await client.post(f'/api/cases/{case_id}/documents', headers=headers,
                                 files={'file': ('record.txt', source, 'text/plain')})
    assert uploaded.status_code == 201, uploaded.text
    review = (await client.get(f'/api/cases/{case_id}/review', headers=headers)).json()
    return case_id, headers, source, review


def selected_ids(review):
    return [c['id'] for c in review['candidate_codes'] if c['code'] in {'K35.80', '99285'}]


@pytest.mark.asyncio
async def test_review_and_signed_export(client, seed_data):
    case_id, headers, _, review = await prepared_case(client, seed_data)
    for fact in review['clinical_facts']:
        ev = fact['evidence']
        page = next(p for d in review['documents'] if d['id'] == ev['document_id'] for p in d['pages'] if p['page_number'] == ev['page_number'])
        assert page['ocr_text'][ev['start_offset']:ev['end_offset']] == ev['quoted_text']
        assert ev['bbox'] is None
    decision = {'decision': 'accept', 'version': review['version'], 'approved_candidate_ids': selected_ids(review)}
    accepted = await client.post(f'/api/cases/{case_id}/decisions', headers=headers, json=decision)
    assert accepted.status_code == 201, accepted.text
    stale = await client.post(f'/api/cases/{case_id}/decisions', headers=headers, json=decision)
    assert stale.status_code == 409
    exported = await client.post(f'/api/cases/{case_id}/exports', headers=headers)
    assert exported.status_code == 200, exported.text
    manifest = exported.json()
    signature = manifest.pop('integrity_signature')
    expected = hmac.new(settings.SECRET_KEY.encode(), json.dumps(manifest, sort_keys=True, separators=(',', ':')).encode(), hashlib.sha256).hexdigest()
    assert hmac.compare_digest(signature, expected)
    assert {c['code'] for c in manifest['approved_codes']} == {'K35.80', '99285'}
    assert (await client.post(f'/api/cases/{case_id}/exports', headers=headers)).status_code == 200
    for format in ('json', 'csv', 'worksheet', 'collaboratemd'):
        assert (await client.get(f'/api/exports/encounter/{case_id}/{format}', headers=headers)).status_code == 200


@pytest.mark.asyncio
async def test_source_change_preserves_history_and_invalidates_approval(client, seed_data, db_session):
    case_id, headers, source, review = await prepared_case(client, seed_data)
    accepted = await client.post(f'/api/cases/{case_id}/decisions', headers=headers, json={
        'decision': 'accept', 'version': review['version'], 'approved_candidate_ids': selected_ids(review),
    })
    assert accepted.status_code == 201
    duplicate = await client.post(f'/api/cases/{case_id}/documents', headers=headers, files={'file': ('same.txt', source, 'text/plain')})
    assert duplicate.json()['pipeline_result']['duplicate']
    assert (await client.get(f'/api/cases/{case_id}/review', headers=headers)).json()['status'] == 'approved'
    amendment = source + b'\nAdditional documentation: Follow-up required.'
    assert (await client.post(f'/api/cases/{case_id}/documents', headers=headers, files={'file': ('amended.txt', amendment, 'text/plain')})).status_code == 201
    latest = (await client.get(f'/api/cases/{case_id}/review', headers=headers)).json()
    assert latest['status'] == 'review_required'
    assert latest['snapshot_id'] != review['snapshot_id']
    assert len(latest['candidate_codes']) == len(review['candidate_codes'])
    assert len((await db_session.execute(select(CaseSnapshot).where(CaseSnapshot.case_id == case_id))).scalars().all()) == 3
    assert (await client.post(f'/api/cases/{case_id}/exports', headers=headers)).status_code == 400
    wrong_ids = await client.post(f'/api/cases/{case_id}/decisions', headers=headers, json={
        'decision': 'accept', 'version': latest['version'], 'approved_candidate_ids': selected_ids(review),
    })
    assert wrong_ids.status_code == 400


@pytest.mark.asyncio
async def test_tampered_evidence_blocks_approval(client, seed_data, db_session):
    case_id, headers, _, review = await prepared_case(client, seed_data)
    candidate = next(c for c in review['candidate_codes'] if c['code'] == 'K35.80')
    evidence = await db_session.get(Evidence, candidate['supporting_evidence'][0]['id'])
    evidence.quoted_text = 'fabricated diagnosis'
    await db_session.commit()
    response = await client.post(f'/api/cases/{case_id}/decisions', headers=headers, json={
        'decision': 'accept', 'version': review['version'], 'approved_candidate_ids': [candidate['id']],
    })
    assert response.status_code == 400


@pytest.mark.asyncio
async def test_all_exports_block_unapproved_case(client, seed_data):
    for format in ('json', 'csv', 'worksheet', 'collaboratemd'):
        response = await client.get(f'/api/exports/encounter/{seed_data["encounter1"].id}/{format}', headers={'Authorization': 'Bearer ' + seed_data['tokens']['reviewer_org1']})
        assert response.status_code == 400


@pytest.mark.asyncio
async def test_collaboratemd_access_checks(client, seed_data, db_session):
    case_id, headers, source, review = await prepared_case(client, seed_data)
    path = f'/api/exports/encounter/{case_id}/collaboratemd'
    assert (await client.get(path)).status_code == 401
    assert (await client.get(path, headers={'Authorization': 'Bearer ' + seed_data['tokens']['coder_org2']})).status_code == 404
    assert (await client.get(path, headers=headers)).status_code == 400


@pytest.mark.asyncio
async def test_reviewed_er_split_populates_excel_and_manifest(client, seed_data):
    case_id, headers, source, _ = await prepared_case(client, seed_data)
    basis = 'Facility resources: ED nursing assessment and monitoring performed; level assessed using SYNTH-HOSPITAL-POLICY-v1.'
    uploaded = await client.post(f'/api/cases/{case_id}/documents', headers=headers,
        files={'file': ('facility.txt', source + b'\n' + basis.encode(), 'text/plain')})
    assert uploaded.status_code == 201
    review = (await client.get(f'/api/cases/{case_id}/review', headers=headers)).json()
    dx = next(c for c in review['candidate_codes'] if c['code'] == 'K35.80')
    em = next(c for c in review['candidate_codes'] if c['code'] == '99285')
    evidence = em['supporting_evidence'][0]
    page = next(p for d in review['documents'] if d['id'] == evidence['document_id'] for p in d['pages'] if basis in p['ocr_text'])
    start = page['ocr_text'].index(basis)
    accepted = await client.post(f'/api/cases/{case_id}/decisions', headers=headers,
        json={'decision': 'accept', 'version': review['version'], 'approved_candidate_ids': [dx['id'], em['id']]})
    assert accepted.status_code == 201
    line = {'claim_type': 'professional', 'candidate_id': em['id'], 'service_type': 'EM',
            'units': 1, 'diagnosis_pointers': [1], 'supporting_evidence_ids': [evidence['id']]}
    facility = {**line, 'claim_type': 'facility', 'revenue_category': 'ER_FACILITY',
                'facility_policy_reference': 'SYNTH-HOSPITAL-POLICY-v1',
                'facility_basis': {'page_number': page['page_number'], 'start_offset': start,
                                   'end_offset': start + len(basis), 'quoted_text': basis}}
    data = {'version': accepted.json()['version'], 'snapshot_id': review['snapshot_id'],
            'diagnosis_candidate_ids': [dx['id']], 'notes': 'Synthetic resource review.', 'lines': [line, facility]}
    path = f'/api/cases/{case_id}/split-billing'
    coder = {'Authorization': 'Bearer ' + seed_data['tokens']['coder_org1']}
    assert (await client.post(path, headers=coder, json=data)).status_code == 403
    foreign = {'Authorization': 'Bearer ' + seed_data['tokens']['reviewer_org1']}
    assert (await client.post('/api/cases/enc-2/split-billing', headers=foreign, json=data)).status_code == 404
    bad = {**data, 'lines': [line, {**facility, 'facility_basis': {**facility['facility_basis'], 'quoted_text': 'Fabricated resources'}}]}
    assert (await client.post(path, headers=headers, json=bad)).status_code == 400
    split = await client.post(path, headers=headers, json=data)
    assert split.status_code == 200, split.text
    assert split.json()['ub04_lines'][0]['rev_code'] == '0450'
    assert split.json()['cms1500_lines'][0]['primary_dx'] == 'K35.80'
    assert (await client.post(path, headers=headers, json=data)).status_code == 409
    workbook = await client.get(f'/api/exports/encounter/{case_id}/collaboratemd', headers=headers)
    assert workbook.status_code == 200, workbook.text
    wb = load_workbook(BytesIO(workbook.content))
    assert wb.worksheets[0]['D7'].value == 'K35.80'
    assert wb.worksheets[0]['J7'].value == 1
    assert wb.worksheets[1]['C7'].value == '0450'
    assert wb.worksheets[1]['C7'].data_type == 's'
    assert basis in wb.worksheets[1]['J7'].value
    manifest = await client.post(f'/api/cases/{case_id}/exports', headers=headers)
    assert manifest.status_code == 200, manifest.text
    assert manifest.json()['split_billing']['billing_review_id'] == split.json()['billing_review_id']
    # A new coding approval on the same source invalidates prior billing selection.
    latest = (await client.get(f'/api/cases/{case_id}/review', headers=headers)).json()
    edited = await client.post(f'/api/cases/{case_id}/decisions', headers=headers,
        json={'decision': 'edit', 'version': latest['version'], 'approved_candidate_ids': []})
    assert edited.status_code == 201
    approved = await client.post(f'/api/cases/{case_id}/decisions', headers=headers,
        json={'decision': 'accept', 'version': edited.json()['version'], 'approved_candidate_ids': [dx['id'], em['id']]})
    assert approved.status_code == 201
    assert (await client.get(f'/api/exports/encounter/{case_id}/collaboratemd', headers=headers)).status_code == 409


@pytest.mark.asyncio
async def test_collaboratemd_workbook_uses_reviewed_data_and_audits_hash(client, seed_data, db_session):
    case_id, headers, source, review = await prepared_case(client, seed_data)
    path = f'/api/exports/encounter/{case_id}/collaboratemd'
    assert (await client.post(f'/api/cases/{case_id}/decisions', headers=headers, json={
        'decision': 'accept', 'version': review['version'], 'approved_candidate_ids': selected_ids(review),
    })).status_code == 201
    response = await client.get(path, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert response.headers['cache-control'] == 'no-store'
    assert response.headers['content-disposition'].endswith('.xlsx"')
    assert int(response.headers['content-length']) == len(response.content)
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ['CMS-1500 (Professional)', 'UB-04 (Institutional)']
    ws = wb.worksheets[0]
    assert ws['B7'].value.date().isoformat() == '2026-09-22'
    assert ws['C7'].value == '23'
    assert ws['F7'].value == '99285'
    for column in ('D', 'E', 'G', 'H', 'I', 'J', 'K'):
        assert ws[f'{column}7'].value is None
    assert 'Document ' in ws['L7'].value
    text = ' '.join(str(cell.value) for sheet in wb for row in sheet for cell in row if cell.value)
    assert 'K35.80' in text
    assert 'S93.401A' not in text
    assert '99284' not in text
    assert '0450' not in text
    assert wb.worksheets[1].max_row == 6
    assert 'Facility lines unavailable' in wb.worksheets[1]['A4'].value
    event = (await db_session.execute(select(AuditEvent).where(AuditEvent.entity_id == case_id,
                       AuditEvent.event_type == 'encounter_exported_xlsx'))).scalar_one()
    assert event.actor_id == seed_data['user_reviewer'].id
    assert event.metadata_json['sha256'] == hashlib.sha256(response.content).hexdigest()
    assert event.metadata_json['snapshot_id'] == review['snapshot_id']
    assert set(event.metadata_json['candidate_ids']) == set(selected_ids(review))
    # A fresh source invalidates this export just like the existing formats.
    assert (await client.post(f'/api/cases/{case_id}/documents', headers=headers,
        files={'file': ('amended.txt', source + b'\nFollow-up advised.', 'text/plain')})).status_code == 201
    assert (await client.get(path, headers=headers)).status_code == 400


@pytest.mark.asyncio
async def test_strict_case_inputs(client, seed_data):
    headers = {'Authorization': 'Bearer ' + seed_data['tokens']['reviewer_org1']}
    for overrides in ({'service_date': '2026-02-30'}, {'coding_domain': 'inpatient'}, {'patient_reference': '  '}):
        response = await client.post('/api/cases', headers=headers, json={
            'patient_reference': 'SYNTH', 'service_date': '2026-09-22', **overrides,
        })
        assert response.status_code == 422
