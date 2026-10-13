from io import BytesIO
import hashlib
import json
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from services.api.db.models import ClinicalFact, CodeCandidate

BH = '''Date of Service: 2026-09-22
Service kind: PSYCHOTHERAPY
Psychotherapy duration: 45 minutes
Patient present: yes
Telehealth: no
Separate medical E/M: yes
Separate E/M MDM: LOW
Establishment confirmed: yes
Signed by: Synthetic Reviewer
'''
DME = '''Date of Service: 2026-09-22
Equipment: CPAP
Order type: RENTAL
Signed CMN on file: yes
Sleep study criteria met: yes
Signed by: Synthetic Reviewer
'''


async def prepared(client, seed_data, domain, text, pos):
    headers = {'Authorization': 'Bearer ' + seed_data['tokens']['reviewer_org1']}
    response = await client.post('/api/cases', headers=headers, json={
        'patient_reference': 'SYNTH-V02', 'service_date': '2026-09-22', 'coding_domain': domain, 'place_of_service': pos})
    assert response.status_code == 201, response.text
    cid = response.json()['id']
    response = await client.post(f'/api/cases/{cid}/documents', headers=headers,
        files={'file': ('synthetic.txt', text.encode(), 'text/plain')})
    assert response.status_code == 201, response.text
    review = (await client.get(f'/api/cases/{cid}/review', headers=headers)).json()
    return cid, headers, review


@pytest.mark.asyncio
async def test_behavioral_addon_cannot_be_approved_alone(client, seed_data):
    cid, headers, review = await prepared(client, seed_data, 'behavioral_health', BH, '11')
    assert {c['code'] for c in review['candidate_codes']} == {'99213','90836'}
    addon = next(c for c in review['candidate_codes'] if c['code'] == '90836')
    response = await client.post(f'/api/cases/{cid}/decisions', headers=headers, json={
        'decision':'accept', 'version':review['version'], 'approved_candidate_ids':[addon['id']]})
    assert response.status_code == 400
    assert 'add-on' in response.json()['error']['message']


@pytest.mark.asyncio
async def test_dme_explicit_modifier_acceptance_and_stamped_export(client, seed_data, db_session):
    cid, headers, review = await prepared(client, seed_data, 'dme', DME, '12')
    assert review['specialty_status']['status'] == 'Available'
    c = review['candidate_codes'][0]
    assert c['code'] == 'E0601' and c['proposed_modifiers'] == ['RR','KX']
    decision = {'decision':'accept', 'version':review['version'], 'approved_candidate_ids':[c['id']]}
    assert (await client.post(f'/api/cases/{cid}/decisions', headers=headers, json=decision)).status_code == 400
    decision['accepted_modifiers'] = {c['id']: ['RR','KX']}
    response = await client.post(f'/api/cases/{cid}/decisions', headers=headers, json=decision)
    assert response.status_code == 201, response.text
    response = await client.get(f'/api/exports/encounter/{cid}/collaboratemd', headers=headers)
    assert response.status_code == 200, response.text
    wb = load_workbook(BytesIO(response.content))
    assert [wb.active.cell(6,col).value for col in range(1,13)] == ['Line #','DOS','POS','Primary DX','Secondary DXs','CPT/HCPCS','Mod 1','Mod 2','DX Pointer','Units','Charge','Rationale']
    assert (wb.active['F7'].value,wb.active['G7'].value,wb.active['H7'].value) == ('E0601','RR','KX')
    assert wb.active['K7'].value is None
    canonical = wb.custom_doc_props['RemitVistaManifestJSON'].value
    digest = hashlib.sha256(canonical.encode()).hexdigest()
    assert wb.custom_doc_props['RemitVistaManifestSHA256'].value == digest
    assert digest in wb.active['A5'].value
    assert json.loads(canonical)['snapshot_id'] == review['snapshot_id']
    # Mutating persisted proposals cannot change accepted modifiers on export.
    candidate = await db_session.get(CodeCandidate, c['id'])
    candidate.modifiers_json = ['NU','KX']
    await db_session.commit()
    assert (await client.get(f'/api/exports/encounter/{cid}/collaboratemd', headers=headers)).status_code == 400


@pytest.mark.asyncio
@pytest.mark.parametrize('domain,text,pos', [
    ('behavioral_health',BH.replace('Psychotherapy duration: 45 minutes\n',''),'11'),
    ('behavioral_health',BH.replace('45 minutes','12 minutes'),'11'),
    ('dme','Equipment: semi-electric hospital bed with mattress\nOrder type: PURCHASE\nSigned CMN on file: no','12')])
async def test_omission_and_documentation_blocks_are_visible(client, seed_data, domain, text, pos):
    cid, headers, review = await prepared(client, seed_data, domain, text, pos)
    assert review['specialty_status']['status'] == 'Unavailable'
    assert review['candidate_codes'] == []
    assert any(f['severity'] == 'BLOCK' for f in review['validation_findings'])
    assert (await client.get(f'/api/exports/encounter/{cid}/collaboratemd', headers=headers)).status_code == 400


@pytest.mark.asyncio
async def test_specialty_facts_are_reextracted_before_approval(client, seed_data, db_session):
    cid, headers, review = await prepared(client, seed_data, 'behavioral_health', BH, '11')
    fact = (await db_session.execute(select(ClinicalFact).where(ClinicalFact.encounter_id == cid,
            ClinicalFact.fact_type == 'behavioral_health'))).scalars().all()
    minutes = next(f for f in fact if 'session_duration_minutes' in f.value_json)
    minutes.value_json = {'session_duration_minutes':60}
    await db_session.commit()
    response = await client.post(f'/api/cases/{cid}/decisions', headers=headers, json={
        'decision':'accept', 'version':review['version'], 'approved_candidate_ids':[c['id'] for c in review['candidate_codes']]})
    assert response.status_code == 400
    assert 'independently extracted' in response.json()['error']['message']


@pytest.mark.asyncio
async def test_telehealth_location_and_modifiers_survive_export(client, seed_data):
    text = BH.replace('Telehealth: no','Telehealth: yes\nTelehealth modality: AUDIO_VIDEO\nTelehealth patient location: HOME')
    cid, headers, review = await prepared(client, seed_data, 'behavioral_health', text, '10')
    decision = {'decision':'accept','version':review['version'],'approved_candidate_ids':[c['id'] for c in review['candidate_codes']],
                'accepted_modifiers':{c['id']:c['proposed_modifiers'] for c in review['candidate_codes']}}
    assert (await client.post(f'/api/cases/{cid}/decisions', headers=headers, json=decision)).status_code == 201
    response = await client.get(f'/api/exports/encounter/{cid}/collaboratemd', headers=headers)
    assert response.status_code == 200, response.text
    ws = load_workbook(BytesIO(response.content)).active
    assert ws['C7'].value == '10' and ws['G7'].value == '95'
