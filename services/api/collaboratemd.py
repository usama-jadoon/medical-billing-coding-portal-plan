"""Manual-entry worksheets. Vendor import compatibility is not established."""
from datetime import date
from io import BytesIO
import math
import hashlib
import json
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.packaging.custom import StringProperty


class _InMemoryExcelWriter(ExcelWriter):
    """openpyxl 3.1.5 normally writes worksheet XML to temporary disk files.

    Our cell-only worksheets instead serialize directly to memory. Keep the
    dependency pinned and rerun the no-temporary-files test before upgrading.
    """
    def write_worksheet(self, ws):
        with BytesIO() as xml:
            writer = WorksheetWriter(ws, out=xml)
            writer.write()
            ws._rels = writer._rels
            self._archive.writestr(ws.path[1:], xml.getvalue())
        self.manifest.append(ws)


PROFESSIONAL_HEADERS = [
    "Line #", "DOS", "POS", "Primary DX", "Secondary DXs", "CPT/HCPCS", "Mod 1", "Mod 2",
    "DX Pointer", "Units", "Charge", "Rationale",
]
FACILITY_HEADERS = [
    "Line #", "DOS", "Rev Code", "Rev Description", "HCPCS/CPT", "Modifier", "Units", "Primary DX", "Pointer", "Justification",
]


def _write(cell, value):
    # Preserve identifiers and hostile source text as literal strings, never formulas.
    if isinstance(value, str):
        if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
            raise ValueError("Export text exceeds Excel limits or contains unsupported control characters.")
        cell.value = value
        cell.data_type = "s"
        cell.number_format = "@"
    else:
        cell.value = value
        if isinstance(value, date):
            cell.number_format = "mm/dd/yyyy"
    cell.font = Font(name="Calibri", size=10)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _header(ws, row, headers, color):
    for col, label in enumerate(headers, 1):
        cell = ws.cell(row, col)
        _write(cell, label)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="CBD5E1"))
    ws.row_dimensions[row].height = 48


def _sheet(wb, title, headers, color, payload, note):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = True
    for row, text in enumerate([
        "CollaborateMD manual-entry worksheet",
        f"Encounter: {payload['encounter_id']}    Patient reference: {payload['patient_reference']}",
        f"Approved snapshot: {payload['snapshot_id']}", note,
    ], 1):
        _write(ws.cell(row, 1), text)
        # Context is separate from the working table, so it can span the page.
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        ws.row_dimensions[row].height = 30 if row == 4 else 22
    ws.cell(1, 1).font = Font(name="Calibri", size=14, bold=True, color=color)
    _header(ws, 6, headers, color)
    ws.freeze_panes = "D7"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"
    return ws


def _rows(ws, rows, keys):
    for row_idx, item in enumerate(rows, 7):
        values = [row_idx - 6] + [item.get(key) for key in keys]
        for col, value in enumerate(values, 1):
            _write(ws.cell(row_idx, col), value)
        ws.row_dimensions[row_idx].height = 30
    ws.auto_filter.ref = f"A6:{get_column_letter(len(keys) + 1)}{max(6, 6 + len(rows))}"


def _fit(ws, table_end, diagnosis_column=False):
    for col in range(1, ws.max_column + 1):
        lengths = [len(str(ws.cell(row, col).value or "")) for row in range(6, table_end + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(45, max(12, max(lengths, default=0) + 3))
    if diagnosis_column:
        ws.column_dimensions["C"].width = 45
    # Fit wrapped table text; never silently clip source evidence at Excel's row limit.
    for row in range(7, ws.max_row + 1):
        lines = max((sum(max(1, math.ceil(len(part) / max(8, ws.column_dimensions[get_column_letter(cell.column)].width - 3)))
                         for part in str(cell.value or "").split("\n"))
                     for cell in ws[row]), default=1)
        height = max(30, lines * 15)
        if height > 409:
            raise ValueError("Source evidence is too long to display in a worksheet row; use the JSON export.")
        ws.row_dimensions[row].height = height
    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def create_collaboratemd_workbook(encounter_data: dict) -> BytesIO:
    """Build explicit lines only. No invented diagnoses, units or facility charges."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = _sheet(wb, "CMS-1500 (Professional)", PROFESSIONAL_HEADERS, "1E3A8A", encounter_data,
                "Manual entry only. Full payer/NCCI policy verification is unavailable. " +
                ("Claim fields use explicit billing review selections." if encounter_data.get('billing_review_id') else
                 "Blank claim fields need review; diagnosis order and pointers are not assigned."))
    lines = encounter_data.get("cms1500_lines", [])
    _rows(ws, lines, ["dos", "pos", "primary_dx", "secondary_dx", "cpt", "mod1", "mod2", "pointer", "units", "charge", "justification"])
    for row in range(7, 7 + len(lines)):
        ws.cell(row, 11).number_format = '"$"#,##0.00'
    table_end = 6 + len(lines)
    diagnosis_header = max(8, table_end + 3)
    _header(ws, diagnosis_header, ["Approved diagnosis", "Version", "Evidence (unsequenced)"], "1E3A8A")
    for row, dx in enumerate(encounter_data.get("diagnoses", []), diagnosis_header + 1):
        for col, key in enumerate(["code", "version", "evidence"], 1):
            _write(ws.cell(row, col), dx[key])
    _fit(ws, max(6, table_end), bool(encounter_data.get("diagnoses")))
    if encounter_data.get("is_facility", False):
        facility_lines = encounter_data.get('ub04_lines', [])
        facility = _sheet(wb, "UB-04 (Institutional)", FACILITY_HEADERS, "581C87", encounter_data,
                          "Explicit facility review selections. Verify hospital and payer policies before claim entry." if facility_lines else
                          "Facility lines unavailable: this case has no separately reviewed revenue codes or facility charges. No professional lines were copied.")
        _rows(facility, facility_lines, ["dos", "rev_code", "rev_desc", "cpt", "mod", "units", "primary_dx", "pointer", "justification"])
        _fit(facility, max(6, 6 + len(facility_lines)))
    canonical = json.dumps(encounter_data, sort_keys=True, separators=(',', ':'), default=lambda value: value.isoformat() if isinstance(value, date) else str(value))
    stamp = hashlib.sha256(canonical.encode('utf-8')).hexdigest()
    wb.custom_doc_props.append(StringProperty(name='RemitVistaManifestJSON', value=canonical))
    wb.custom_doc_props.append(StringProperty(name='RemitVistaManifestSHA256', value=stamp))
    for sheet in wb.worksheets:
        _write(sheet.cell(5, 1), 'Manifest SHA-256: ' + stamp)
        sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=sheet.max_column)
        sheet.row_dimensions[5].height = 22
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED, allowZip64=True) as archive:
        _InMemoryExcelWriter(wb, archive).write_data()
    output.seek(0)
    return output
